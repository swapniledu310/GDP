"""Complete data-extraction + Excel dashboard section for the
List of countries by GDP (nominal) data set.

Source : https://en.wikipedia.org/wiki/List_of_countries_by_GDP_(nominal)

Structure of this file mirrors IPO_Index_GMP_Full_Colab.py :

    1  scrape          -> df, validation_df, website_total
    2  data types / helper columns
    3  six pivot analyses
    4  Excel workbook  -> Data sheet + Excel Table + Pivot sheet + Dashboard
    5  final count validation (Website = DataFrame = Excel)
    6  download
"""

import importlib.util
import re
import site
import subprocess
import sys
import time

# Install only packages that are missing. This keeps the script runnable as a
# single Google Colab cell as well as a normal Python file.
_requirements = {
    "numpy": "numpy",
    "pandas": "pandas",
    "requests": "requests",
    "bs4": "beautifulsoup4",
    "openpyxl": "openpyxl",
}
_missing = [package for module, package in _requirements.items() if importlib.util.find_spec(module) is None]
if _missing:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "-q", *_missing])
    site.addsitedir(site.getusersitepackages())

import numpy as np
import pandas as pd
import requests
from bs4 import BeautifulSoup


SOURCE_URL = (
    "https://en.wikipedia.org/wiki/List_of_countries_by_GDP_%28nominal%29"
)

# Order of the three estimate blocks on the Wikipedia table
SOURCE_NAMES = [
    "IMF",
    "World Bank",
    "United Nations",
]


# =====================================================================
# 0. SMALL TEXT HELPERS
# =====================================================================

def _clean(value):
    text = re.sub(r"\[[^\]]*\]", " ", str(value or ""))      # drop [1], [n 1]
    text = text.replace("\xa0", " ")
    return re.sub(r"\s+", " ", text).strip()


def _number(value):
    """'32,383,920' -> 32383920.0 ; '—' / '' -> NaN"""
    text = _clean(value).replace(",", "")
    if text in {"", "-", "—", "–", "N/A", "n/a"}:
        return np.nan
    match = re.search(r"[-+]?\d+(?:\.\d+)?", text)
    return float(match.group()) if match else np.nan


def _is_year(value):
    text = _clean(value)
    return bool(re.fullmatch(r"(19|20|21)\d{2}", text))


def _looks_numeric(value):
    text = _clean(value).replace(",", "")
    return bool(re.fullmatch(r"[-+]?\d+(?:\.\d+)?", text))


# =====================================================================
# 1. SCRAPE WIKIPEDIA
# =====================================================================

def _pick_gdp_table(soup):
    """Return the wikitable that holds the IMF / World Bank / UN estimates."""
    for table in soup.find_all("table", class_="wikitable"):
        head = _clean(table.get_text(" ", strip=True))[:600].upper()
        if "IMF" in head and "WORLD BANK" in head and "UNITED NATIONS" in head:
            return table
    return None


def _header_years(table):
    """Fallback years taken from the header, e.g. 'IMF (2026)'."""
    head_text = _clean(" ".join(th.get_text(" ", strip=True) for th in table.find_all("th")))
    years = {}
    for name in SOURCE_NAMES:
        match = re.search(rf"{re.escape(name)}\D{{0,12}}((?:19|20|21)\d{{2}})", head_text, flags=re.I)
        years[name] = int(match.group(1)) if match else np.nan
    return years


def _row_texts(cells):
    """Cleaned cell texts with blank cells and a leading rank cell removed."""
    texts = [_clean(cell.get_text(" ", strip=True)) for cell in cells]
    texts = [t for t in texts if t not in {"", "•"}]

    # Some renderings carry a rank column ("1", "2", ...) before the country.
    if (
        len(texts) > 1
        and re.fullmatch(r"\d{1,3}", texts[0])
        and not _looks_numeric(texts[1])
    ):
        texts = texts[1:]

    return texts


def _parse_row(cells, fallback_years, position):
    """
    Wikipedia ships two shapes of this table:

        Country | UN region | IMF | Year | WB | Year | UN | Year
        Country | IMF | WB | UN                    (year only in header)

    Both are handled: the country is the first cell, an optional
    non-numeric second cell is the UN region, then estimates and
    optional 4-digit years are consumed in order.
    """
    texts = _row_texts(cells)

    if not texts:
        raise ValueError(f"Row {position} is empty")

    country = texts[0]
    rest = texts[1:]

    region = ""
    if rest and not _looks_numeric(rest[0]) and not _is_year(rest[0]):
        region = rest[0] if rest[0] not in {"-", "—", "–"} else ""
        rest = rest[1:]

    record = {"Country": country, "UN Region": region or "Not classified"}

    index = 0
    for name in SOURCE_NAMES:
        estimate = np.nan
        year = fallback_years.get(name, np.nan)

        if index < len(rest):
            token = rest[index]
            if _is_year(token) and "," not in token:
                # stray year with no estimate -> no value for this source
                pass
            else:
                estimate = _number(token)
                index += 1

        if index < len(rest) and _is_year(rest[index]) and "," not in rest[index]:
            year = int(_clean(rest[index]))
            index += 1

        record[f"{name} Estimate"] = estimate
        record[f"{name} Year"] = year

    if all(pd.isna(record[f"{name} Estimate"]) for name in SOURCE_NAMES):
        raise ValueError(f"No GDP value found for '{country}'")

    record["_Website Position"] = position
    record["Source URL"] = SOURCE_URL
    return record


def scrape_gdp_table():
    """Download and parse every country row of the Wikipedia GDP table."""
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126 Safari/537.36"
        ),
        "Accept-Language": "en-US,en;q=0.9",
    }

    last_error = None
    for attempt in range(1, 4):
        try:
            response = requests.get(SOURCE_URL, headers=headers, timeout=45)
            response.raise_for_status()
            if "WORLD BANK" not in response.text.upper():
                raise RuntimeError("The page loaded, but the GDP table was not present in its HTML.")
            break
        except Exception as exc:                                 # noqa: BLE001
            last_error = exc
            if attempt == 3:
                raise RuntimeError(f"Unable to load {SOURCE_URL}: {exc}") from last_error
            time.sleep(attempt * 2)

    soup = BeautifulSoup(response.text, "html.parser")
    table = _pick_gdp_table(soup)

    if table is None:
        raise RuntimeError(
            "The IMF / World Bank / United Nations table was not found. "
            "The article layout may have changed; open SOURCE_URL once and rerun."
        )

    fallback_years = _header_years(table)

    body_rows = []
    for tr in table.find_all("tr"):
        cells = tr.find_all(["td", "th"])
        if not cells:
            continue
        if all(cell.name == "th" for cell in cells):            # header rows
            continue
        body_rows.append(cells)

    if not body_rows:
        raise RuntimeError("The GDP table was found, but it contains no data rows.")

    world_row = None
    records, errors, position = [], [], 0

    for cells in body_rows:
        texts = _row_texts(cells)
        first = texts[0] if texts else ""

        if first.upper() == "WORLD":
            try:
                world_row = _parse_row(cells, fallback_years, 0)
            except Exception:                                   # noqa: BLE001
                world_row = None
            continue

        position += 1
        try:
            records.append(_parse_row(cells, fallback_years, position))
        except Exception as exc:                                # noqa: BLE001
            errors.append({"Website Position": position, "Result": "ERROR", "Message": str(exc)})

    if not records:
        raise RuntimeError("Table rows were found, but none could be parsed.")

    website_total = position
    df = pd.DataFrame(records)

    good = pd.DataFrame({
        "Website Position": df["_Website Position"],
        "Country": df["Country"],
        "Result": "OK",
        "Message": "Parsed successfully",
    })
    validation_df = pd.concat([good, pd.DataFrame(errors)], ignore_index=True, sort=False)

    if len(df) != website_total:
        failed = website_total - len(df)
        examples = "; ".join(item["Message"] for item in errors[:3])
        raise RuntimeError(f"Parsed {len(df)} of {website_total} rows; {failed} failed. {examples}")

    print(f"Source: {SOURCE_URL}")
    print(f"Country rows extracted: {len(df)}")
    return df, validation_df, website_total, world_row


# Create the objects required by the Excel/dashboard section below.
df, validation_df, website_total, world_row = scrape_gdp_table()


# =====================================================================
# EXCEL OUTPUT
# Pivot Analysis + Interactive Dashboard
# =====================================================================

# =====================================================================
# 1. EXTRA IMPORTS
# =====================================================================

import numpy as np
import pandas as pd

from openpyxl import load_workbook

from openpyxl.styles import (
    Font,
    Alignment,
    PatternFill,
    Border,
    Side
)

from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

from openpyxl.worksheet.table import (
    Table,
    TableStyleInfo
)

from openpyxl.worksheet.datavalidation import (
    DataValidation
)

from openpyxl.chart import (
    BarChart,
    LineChart,
    DoughnutChart,
    Reference
)

try:
    from google.colab import files
except ImportError:
    class _LocalFiles:
        @staticmethod
        def download(path):
            print(f"Google Colab was not detected. File saved locally: {path}")

    files = _LocalFiles()


# =====================================================================
# 2. SETTINGS
# =====================================================================

OUTPUT_FILE = "GDP_Nominal_Dashboard.xlsx"

CSV_FILE = "GDP_Nominal_Data.csv"

DATA_SHEET = "GDP Data"

PIVOT_SHEET = "Pivot Analysis"

DASHBOARD_SHEET = "Dashboard"

VALIDATION_SHEET = "Validation"

LISTS_SHEET = "_Lists"

EXCEL_TABLE_NAME = "GDP_DATA"


# =====================================================================
# 3. DATA TYPES
# =====================================================================

# -------------------------------------------------------------
# Numeric estimate columns (million US$)
# -------------------------------------------------------------

ESTIMATE_COLUMNS = [
    "IMF Estimate",
    "World Bank Estimate",
    "United Nations Estimate"
]


for column in ESTIMATE_COLUMNS:

    df[column] = pd.to_numeric(
        df[column],
        errors="coerce"
    )


# -------------------------------------------------------------
# Year columns
# -------------------------------------------------------------

YEAR_COLUMNS = [
    "IMF Year",
    "World Bank Year",
    "United Nations Year"
]


for column in YEAR_COLUMNS:

    df[column] = pd.to_numeric(
        df[column],
        errors="coerce"
    )


# -------------------------------------------------------------
# Text columns
# -------------------------------------------------------------

for column in [
    "Country",
    "UN Region"
]:

    df[column] = (
        df[column]
        .astype(str)
        .str.strip()
    )


# =====================================================================
# 4. HELPER COLUMNS
# =====================================================================

# -------------------------------------------------------------
# GDP Bucket used by the Dashboard size filter
# -------------------------------------------------------------

BUCKET_LABELS = [

    "Under $10 bn",

    "$10 bn - $100 bn",

    "$100 bn - $1 tn",

    "Above $1 tn"
]


df["GDP Bucket"] = pd.cut(

    df["IMF Estimate"],

    bins=[
        -np.inf,
        10_000,
        100_000,
        1_000_000,
        np.inf
    ],

    labels=BUCKET_LABELS

).astype(str)


df["GDP Bucket"] = (
    df["GDP Bucket"]
    .replace(
        "nan",
        "Not available"
    )
)


# -------------------------------------------------------------
# Share of the world total (IMF basis)
# -------------------------------------------------------------

if (
    world_row
    and not pd.isna(
        world_row.get("IMF Estimate")
    )
):

    world_imf_total = float(
        world_row["IMF Estimate"]
    )

else:

    world_imf_total = float(
        df["IMF Estimate"].sum()
    )


df["Share of World"] = (
    df["IMF Estimate"]
    / world_imf_total
)


# -------------------------------------------------------------
# How many of the three sources report this country
# -------------------------------------------------------------

df["Sources Available"] = (
    df[ESTIMATE_COLUMNS]
    .notna()
    .sum(axis=1)
)


# -------------------------------------------------------------
# Spread between the highest and lowest of the three estimates
# -------------------------------------------------------------

df["Estimate Spread"] = (
    df[ESTIMATE_COLUMNS].max(axis=1)
    - df[ESTIMATE_COLUMNS].min(axis=1)
)


# =====================================================================
# 5. SORT BY IMF ESTIMATE (LARGEST FIRST)
# =====================================================================

df = df.sort_values(

    by="IMF Estimate",

    ascending=False,

    kind="stable",

    na_position="last"

).reset_index(drop=True)


# -------------------------------------------------------------
# World rank on the IMF estimate
# -------------------------------------------------------------

df["IMF Rank"] = (
    df["IMF Estimate"]
    .rank(
        ascending=False,
        method="min"
    )
)


# =====================================================================
# 6. REMOVE INTERNAL WEBSITE POSITION
# =====================================================================

df = df.drop(

    columns=[
        "_Website Position"
    ],

    errors="ignore"
)


# =====================================================================
# 7. SERIAL NUMBER AS FIRST COLUMN
# =====================================================================

if "S.No." in df.columns:

    df = df.drop(
        columns=["S.No."]
    )


df.insert(
    0,
    "S.No.",
    range(
        1,
        len(df) + 1
    )
)


# =====================================================================
# 8. FINAL DATA COLUMN ORDER
# =====================================================================

VISIBLE_COLUMNS = [

    "S.No.",

    "Country",

    "UN Region",

    "IMF Rank",

    "IMF Estimate",

    "IMF Year",

    "World Bank Estimate",

    "World Bank Year",

    "United Nations Estimate",

    "United Nations Year",

    "Share of World",

    "Estimate Spread"
]


# Helper columns stay at end and will be hidden in Excel

FINAL_COLUMNS = (

    VISIBLE_COLUMNS

    + [

        "GDP Bucket",

        "Sources Available"

    ]

)


df = df[
    FINAL_COLUMNS
]


# =====================================================================
# 9. WEBSITE COUNT VALIDATION
# =====================================================================

print(
    "\n=============================================="
)

print(
    "PRE-EXCEL COUNT VALIDATION"
)

print(
    "=============================================="
)


print(
    "Website total :",
    website_total
)


print(
    "DataFrame rows:",
    len(df)
)


if len(df) != website_total:

    raise Exception(

        f"COUNT MISMATCH\n"
        f"Website = {website_total}\n"
        f"DataFrame = {len(df)}"
    )


print(
    "✅ WEBSITE TOTAL = DATAFRAME ROWS"
)


# =====================================================================
# 10. CREATE SIX PIVOT-STYLE ANALYSES
# =====================================================================

analysis_df = df.copy()


REGION_ORDER = sorted(

    analysis_df["UN Region"]
    .dropna()
    .unique()
    .tolist()
)


# ---------------------------------------------------------------------
# Analysis 1
# Region Summary
# ---------------------------------------------------------------------

pivot_region = pd.pivot_table(

    analysis_df,

    index="UN Region",

    values=[
        "Country",
        "IMF Estimate",
        "Share of World"
    ],

    aggfunc={

        "Country": "count",

        "IMF Estimate": ["sum", "mean"],

        "Share of World": "sum"

    },

    fill_value=0

)


pivot_region.columns = [

    "Countries",

    "Avg GDP (US$ mn)",

    "Total GDP (US$ mn)",

    "Share of World"
]


pivot_region = (
    pivot_region
    .reset_index()
    .sort_values(
        "Total GDP (US$ mn)",
        ascending=False
    )
)


# ---------------------------------------------------------------------
# Analysis 2
# GDP Bucket Summary
# ---------------------------------------------------------------------

pivot_bucket = pd.pivot_table(

    analysis_df,

    index="GDP Bucket",

    values=[
        "Country",
        "IMF Estimate",
        "Share of World"
    ],

    aggfunc={

        "Country": "count",

        "IMF Estimate": "sum",

        "Share of World": "sum"

    },

    fill_value=0

).reset_index()


pivot_bucket = pivot_bucket.rename(

    columns={

        "Country":
            "Countries",

        "IMF Estimate":
            "Total GDP (US$ mn)",

        "Share of World":
            "Share of World"
    }
)


# ---------------------------------------------------------------------
# Analysis 3
# Region x GDP Bucket
# ---------------------------------------------------------------------

pivot_region_bucket = pd.crosstab(

    analysis_df["UN Region"],

    analysis_df["GDP Bucket"]

).reset_index()


pivot_region_bucket["Total"] = (
    pivot_region_bucket
    .select_dtypes(
        include="number"
    )
    .sum(axis=1)
)


# ---------------------------------------------------------------------
# Analysis 4
# Top 15 Economies
# ---------------------------------------------------------------------

pivot_top = (

    analysis_df

    .nlargest(
        15,
        "IMF Estimate"
    )

    [[
        "IMF Rank",
        "Country",
        "UN Region",
        "IMF Estimate",
        "World Bank Estimate",
        "United Nations Estimate",
        "Share of World"
    ]]

    .reset_index(drop=True)
)


# ---------------------------------------------------------------------
# Analysis 5
# Source Comparison
# ---------------------------------------------------------------------

pivot_source = pd.DataFrame({

    "Source": SOURCE_NAMES,

    "Countries Reported": [

        int(analysis_df[f"{name} Estimate"].notna().sum())

        for name in SOURCE_NAMES
    ],

    "Total GDP (US$ mn)": [

        float(analysis_df[f"{name} Estimate"].sum())

        for name in SOURCE_NAMES
    ],

    "Average GDP (US$ mn)": [

        float(analysis_df[f"{name} Estimate"].mean())

        for name in SOURCE_NAMES
    ],

    "Largest Economy": [

        (
            analysis_df
            .loc[
                analysis_df[f"{name} Estimate"].idxmax(),
                "Country"
            ]
            if analysis_df[f"{name} Estimate"].notna().any()
            else ""
        )

        for name in SOURCE_NAMES
    ]
})


# ---------------------------------------------------------------------
# Analysis 6
# Data Coverage by Region
# ---------------------------------------------------------------------

pivot_coverage = (

    analysis_df

    .groupby(
        "UN Region",
        as_index=False
    )

    .agg(

        Countries=(
            "Country",
            "count"
        ),

        IMF_Reported=(
            "IMF Estimate",
            lambda s: int(s.notna().sum())
        ),

        WorldBank_Reported=(
            "World Bank Estimate",
            lambda s: int(s.notna().sum())
        ),

        UN_Reported=(
            "United Nations Estimate",
            lambda s: int(s.notna().sum())
        ),

        Avg_Sources=(
            "Sources Available",
            "mean"
        )
    )
)


pivot_coverage = pivot_coverage.rename(

    columns={

        "IMF_Reported":
            "IMF Reported",

        "WorldBank_Reported":
            "World Bank Reported",

        "UN_Reported":
            "UN Reported",

        "Avg_Sources":
            "Avg Sources"
    }
)


# =====================================================================
# 11. WRITE BASE WORKBOOK
# =====================================================================

# Nullable years are converted to plain ints/None so openpyxl never
# receives a pandas NA value.

excel_df = df.copy()


for column in YEAR_COLUMNS:

    excel_df[column] = (
        excel_df[column]
        .map(
            lambda v: (
                None
                if pd.isna(v)
                else int(v)
            )
        )
    )


excel_df["IMF Rank"] = (
    excel_df["IMF Rank"]
    .map(
        lambda v: (
            None
            if pd.isna(v)
            else int(v)
        )
    )
)


excel_df.to_csv(
    CSV_FILE,
    index=False
)


with pd.ExcelWriter(

    OUTPUT_FILE,

    engine="openpyxl"

) as writer:


    # Main data

    excel_df.to_excel(

        writer,

        sheet_name=DATA_SHEET,

        index=False
    )


    # Validation

    validation_df.to_excel(

        writer,

        sheet_name=VALIDATION_SHEET,

        index=False
    )


    # -------------------------------------------------------------
    # Six analysis tables
    # -------------------------------------------------------------

    pivot_region.to_excel(

        writer,

        sheet_name=PIVOT_SHEET,

        startrow=2,

        startcol=0,

        index=False
    )


    pivot_bucket.to_excel(

        writer,

        sheet_name=PIVOT_SHEET,

        startrow=2,

        startcol=8,

        index=False
    )


    pivot_region_bucket.to_excel(

        writer,

        sheet_name=PIVOT_SHEET,

        startrow=18,

        startcol=0,

        index=False
    )


    pivot_source.to_excel(

        writer,

        sheet_name=PIVOT_SHEET,

        startrow=18,

        startcol=8,

        index=False
    )


    pivot_top.to_excel(

        writer,

        sheet_name=PIVOT_SHEET,

        startrow=34,

        startcol=0,

        index=False
    )


    pivot_coverage.to_excel(

        writer,

        sheet_name=PIVOT_SHEET,

        startrow=34,

        startcol=9,

        index=False
    )


# =====================================================================
# 12. OPEN WORKBOOK FOR FORMATTING
# =====================================================================

wb = load_workbook(
    OUTPUT_FILE
)


data_ws = wb[
    DATA_SHEET
]


pivot_ws = wb[
    PIVOT_SHEET
]


validation_ws = wb[
    VALIDATION_SHEET
]


# =====================================================================
# 13. CREATE GDP_DATA EXCEL TABLE
# =====================================================================

last_column = get_column_letter(
    data_ws.max_column
)


table_ref = (
    f"A1:{last_column}{data_ws.max_row}"
)


gdp_table = Table(

    displayName=EXCEL_TABLE_NAME,

    ref=table_ref
)


gdp_table.tableStyleInfo = (
    TableStyleInfo(

        name="TableStyleMedium2",

        showFirstColumn=False,

        showLastColumn=False,

        showRowStripes=True,

        showColumnStripes=False
    )
)


data_ws.add_table(
    gdp_table
)


# =====================================================================
# 14. COMMON STYLES
# =====================================================================

dark_blue = "17365D"

blue = "2F65E9"

light_blue = "EAF1FB"

panel_fill = "F4F7FB"

white = "FFFFFF"

green = "16A34A"

red = "DC2626"

orange = "EA580C"

purple = "7C3AED"

gray = "475569"

light_border = Side(

    style="thin",

    color="D9E2F0"
)


# =====================================================================
# 15. FORMAT DATA SHEET
# =====================================================================

data_ws.freeze_panes = "C2"


data_ws.sheet_view.showGridLines = False


header_lookup = {

    data_ws.cell(
        1,
        c
    ).value: c

    for c in range(
        1,
        data_ws.max_column + 1
    )
}


# -------------------------------------------------------------
# Number format helper
# -------------------------------------------------------------

def format_column(
    worksheet,
    header_map,
    column_name,
    number_format
):

    column_number = header_map.get(
        column_name
    )


    if not column_number:

        return


    for row in range(
        2,
        worksheet.max_row + 1
    ):

        worksheet.cell(
            row,
            column_number
        ).number_format = (
            number_format
        )


for column_name in [

    "S.No.",

    "IMF Rank",

    "IMF Estimate",

    "World Bank Estimate",

    "United Nations Estimate",

    "Estimate Spread"

]:

    format_column(
        data_ws,
        header_lookup,
        column_name,
        "#,##0"
    )


for column_name in YEAR_COLUMNS:

    format_column(
        data_ws,
        header_lookup,
        column_name,
        "0"
    )


format_column(
    data_ws,
    header_lookup,
    "Share of World",
    "0.00%"
)


# -------------------------------------------------------------
# Column widths
# -------------------------------------------------------------

widths = {

    "A": 8,

    "B": 30,

    "C": 16,

    "D": 11,

    "E": 18,

    "F": 11,

    "G": 20,

    "H": 13,

    "I": 22,

    "J": 11,

    "K": 15,

    "L": 17,

    "M": 18,

    "N": 16
}


for col, width in widths.items():

    data_ws.column_dimensions[
        col
    ].width = width


# Hide helper columns

for helper in [
    "GDP Bucket",
    "Sources Available"
]:

    col_num = header_lookup[
        helper
    ]

    data_ws.column_dimensions[
        get_column_letter(
            col_num
        )
    ].hidden = True


# =====================================================================
# 16. FORMAT PIVOT ANALYSIS SHEET
# =====================================================================

pivot_ws.sheet_view.showGridLines = False


pivot_ws["A1"] = (
    "GDP (Nominal) - Pivot Analysis"
)


pivot_ws["A1"].font = Font(

    size=18,

    bold=True,

    color=dark_blue
)


titles = {

    "A2":
        "1. Region Analysis",

    "I2":
        "2. GDP Size Bucket Analysis",

    "A18":
        "3. Region × GDP Bucket Analysis",

    "I18":
        "4. Source Comparison",

    "A34":
        "5. Top 15 Economies",

    "J34":
        "6. Data Coverage by Region"
}


for cell, text in titles.items():

    pivot_ws[cell] = text

    pivot_ws[cell].font = Font(

        bold=True,

        size=12,

        color=dark_blue
    )


# =====================================================================
# 17. CREATE EXCEL TABLES FOR SIX ANALYSIS BLOCKS
# =====================================================================

def add_analysis_table(
    ws,
    dataframe,
    start_row,
    start_col,
    table_name
):

    """
    start_row/start_col are zero based pandas positions.
    Excel actual header starts at start_row + 1.
    """

    excel_header_row = (
        start_row + 1
    )

    excel_start_col = (
        start_col + 1
    )

    excel_end_row = (
        excel_header_row
        + len(dataframe)
    )

    excel_end_col = (
        excel_start_col
        + len(dataframe.columns)
        - 1
    )


    ref = (

        f"{get_column_letter(excel_start_col)}"
        f"{excel_header_row}:"

        f"{get_column_letter(excel_end_col)}"
        f"{excel_end_row}"
    )


    tbl = Table(

        displayName=table_name,

        ref=ref
    )


    tbl.tableStyleInfo = (
        TableStyleInfo(

            name="TableStyleMedium4",

            showRowStripes=True,

            showColumnStripes=False
        )
    )


    ws.add_table(
        tbl
    )


add_analysis_table(
    pivot_ws,
    pivot_region,
    2,
    0,
    "PT_Region"
)


add_analysis_table(
    pivot_ws,
    pivot_bucket,
    2,
    8,
    "PT_Bucket"
)


add_analysis_table(
    pivot_ws,
    pivot_region_bucket,
    18,
    0,
    "PT_RegionBucket"
)


add_analysis_table(
    pivot_ws,
    pivot_source,
    18,
    8,
    "PT_Source"
)


add_analysis_table(
    pivot_ws,
    pivot_top,
    34,
    0,
    "PT_TopEconomies"
)


add_analysis_table(
    pivot_ws,
    pivot_coverage,
    34,
    9,
    "PT_Coverage"
)


# Widths

for col in range(
    1,
    18
):

    pivot_ws.column_dimensions[
        get_column_letter(col)
    ].width = 20


# Number formats on the analysis blocks

for row in pivot_ws.iter_rows(
    min_row=3,
    max_row=pivot_ws.max_row,
    min_col=1,
    max_col=pivot_ws.max_column
):

    for cell in row:

        if isinstance(
            cell.value,
            (int, float)
        ):

            if (
                abs(cell.value) <= 1
                and isinstance(
                    cell.value,
                    float
                )
            ):

                cell.number_format = "0.00%"

            else:

                cell.number_format = "#,##0"


# =====================================================================
# 18. CREATE LISTS SHEET FOR FILTERS
# =====================================================================

if LISTS_SHEET in wb.sheetnames:

    del wb[
        LISTS_SHEET
    ]


lists_ws = wb.create_sheet(
    LISTS_SHEET
)


region_options = (
    ["All"]
    + REGION_ORDER
)


for row_num, value in enumerate(
    region_options,
    start=1
):

    lists_ws.cell(
        row_num,
        1
    ).value = value


bucket_options = (
    ["All"]
    + [
        label
        for label in BUCKET_LABELS
        if (df["GDP Bucket"] == label).any()
    ]
)


if (df["GDP Bucket"] == "Not available").any():

    bucket_options.append(
        "Not available"
    )


for row_num, value in enumerate(
    bucket_options,
    start=1
):

    lists_ws.cell(
        row_num,
        2
    ).value = value


lists_ws.sheet_state = "hidden"


# =====================================================================
# 19. CREATE DASHBOARD SHEET
# =====================================================================

if DASHBOARD_SHEET in wb.sheetnames:

    del wb[
        DASHBOARD_SHEET
    ]


dash = wb.create_sheet(
    DASHBOARD_SHEET,
    0
)


dash.sheet_view.showGridLines = False


# -------------------------------------------------------------
# Main title
# -------------------------------------------------------------

dash.merge_cells(
    "A1:F1"
)


dash["A1"] = (
    "GDP (Nominal) Dashboard"
)


dash["A1"].font = Font(

    bold=True,

    size=22,

    color=dark_blue
)


dash["A1"].alignment = Alignment(
    horizontal="left"
)


# =====================================================================
# 20. SLICER-STYLE PANEL
# =====================================================================

for row in range(
    3,
    12
):

    for col in range(
        1,
        7
    ):

        cell = dash.cell(
            row,
            col
        )

        cell.fill = PatternFill(

            "solid",

            fgColor=panel_fill
        )

        cell.border = Border(

            left=light_border,

            right=light_border,

            top=light_border,

            bottom=light_border
        )


dash["A3"] = "FILTERS"


dash["A3"].font = Font(

    size=14,

    bold=True,

    color=dark_blue
)


# -------------------------------------------------------------
# REGION slicer-style control
# -------------------------------------------------------------

dash["A4"] = "REGION"


dash["A4"].font = Font(

    bold=True,

    color=gray
)


dash["C4"] = "All"


dash["C4"].fill = PatternFill(

    "solid",

    fgColor=blue
)


dash["C4"].font = Font(

    bold=True,

    color=white,

    size=12
)


dash["C4"].alignment = Alignment(
    horizontal="center"
)


dash.column_dimensions[
    "C"
].width = 25


region_validation = DataValidation(

    type="list",

    formula1=(
        f"='{LISTS_SHEET}'!"
        f"$A$1:$A${len(region_options)}"
    ),

    allow_blank=False
)


dash.add_data_validation(
    region_validation
)


region_validation.add(
    dash["C4"]
)


# -------------------------------------------------------------
# GDP SIZE slicer-style control
# -------------------------------------------------------------

dash["A7"] = "GDP SIZE"


dash["A7"].font = Font(

    bold=True,

    color=gray
)


dash["C7"] = "All"


dash["C7"].fill = PatternFill(

    "solid",

    fgColor=blue
)


dash["C7"].font = Font(

    bold=True,

    color=white,

    size=12
)


dash["C7"].alignment = Alignment(
    horizontal="center"
)


bucket_validation = DataValidation(

    type="list",

    formula1=(
        f"='{LISTS_SHEET}'!"
        f"$B$1:$B${len(bucket_options)}"
    ),

    allow_blank=False
)


dash.add_data_validation(
    bucket_validation
)


bucket_validation.add(
    dash["C7"]
)


# -------------------------------------------------------------
# Decorative labels
# -------------------------------------------------------------

dash["A9"] = "Region choices:"

dash["B9"] = (
    " | ".join(REGION_ORDER)
)


dash["A10"] = "Size choices:"

dash["B10"] = (
    " | ".join(BUCKET_LABELS)
)


dash["B9"].font = Font(
    size=9,
    color=gray
)


dash["B10"].font = Font(
    size=9,
    color=gray
)


# =====================================================================
# 21. KPI CARDS
# =====================================================================

dash["A13"] = "KEY METRICS"


dash["A13"].font = Font(

    bold=True,

    size=14,

    color=dark_blue
)


# -------------------------------------------------------------
# Helper formulas based on both dashboard filters
#
# Region filter = C4
# Size filter   = C7
# -------------------------------------------------------------

def count_filtered_formula():

    return (

        '=IF(AND($C$4="All",$C$7="All"),'
        'COUNTA(GDP_DATA[Country]),'

        'IF($C$4="All",'
        'COUNTIFS(GDP_DATA[GDP Bucket],$C$7),'

        'IF($C$7="All",'
        'COUNTIFS(GDP_DATA[UN Region],$C$4),'

        'COUNTIFS('
        'GDP_DATA[UN Region],$C$4,'
        'GDP_DATA[GDP Bucket],$C$7))))'
    )


def sum_formula(
    value_column
):

    return (

        '=IF(AND($C$4="All",$C$7="All"),'
        f'SUM(GDP_DATA[{value_column}]),'

        'IF($C$4="All",'
        f'SUMIFS(GDP_DATA[{value_column}],'
        'GDP_DATA[GDP Bucket],$C$7),'

        'IF($C$7="All",'
        f'SUMIFS(GDP_DATA[{value_column}],'
        'GDP_DATA[UN Region],$C$4),'

        f'SUMIFS(GDP_DATA[{value_column}],'
        'GDP_DATA[UN Region],$C$4,'
        'GDP_DATA[GDP Bucket],$C$7))))'
    )


def avg_formula(
    value_column
):

    return (

        '=IFERROR('

        'IF(AND($C$4="All",$C$7="All"),'
        f'AVERAGE(GDP_DATA[{value_column}]),'

        'IF($C$4="All",'
        f'AVERAGEIFS(GDP_DATA[{value_column}],'
        'GDP_DATA[GDP Bucket],$C$7),'

        'IF($C$7="All",'
        f'AVERAGEIFS(GDP_DATA[{value_column}],'
        'GDP_DATA[UN Region],$C$4),'

        f'AVERAGEIFS(GDP_DATA[{value_column}],'
        'GDP_DATA[UN Region],$C$4,'
        'GDP_DATA[GDP Bucket],$C$7))))'

        ',0)'
    )


def max_formula(
    value_column
):

    return (

        '=IFERROR('

        'IF(AND($C$4="All",$C$7="All"),'
        f'MAX(GDP_DATA[{value_column}]),'

        'IF($C$4="All",'
        f'MAXIFS(GDP_DATA[{value_column}],'
        'GDP_DATA[GDP Bucket],$C$7),'

        'IF($C$7="All",'
        f'MAXIFS(GDP_DATA[{value_column}],'
        'GDP_DATA[UN Region],$C$4),'

        f'MAXIFS(GDP_DATA[{value_column}],'
        'GDP_DATA[UN Region],$C$4,'
        'GDP_DATA[GDP Bucket],$C$7))))'

        ',0)'
    )


kpis = [

    (
        "A15",
        "B15",
        "Countries",
        count_filtered_formula(),
        "#,##0"
    ),

    (
        "D15",
        "E15",
        "Total GDP (US$ mn)",
        sum_formula("IMF Estimate"),
        "#,##0"
    ),

    (
        "A18",
        "B18",
        "Average GDP (US$ mn)",
        avg_formula("IMF Estimate"),
        "#,##0"
    ),

    (
        "D18",
        "E18",
        "Largest GDP (US$ mn)",
        max_formula("IMF Estimate"),
        "#,##0"
    ),

    (
        "A21",
        "B21",
        "Share of World",
        sum_formula("Share of World"),
        "0.00%"
    )

]


for (
    label_cell,
    value_cell,
    label,
    formula,
    number_format
) in kpis:


    dash[
        label_cell
    ] = label


    dash[
        label_cell
    ].font = Font(

        bold=True,

        color=gray
    )


    dash[
        value_cell
    ] = formula


    dash[
        value_cell
    ].font = Font(

        bold=True,

        size=16,

        color=dark_blue
    )


    dash[
        value_cell
    ].number_format = (
        number_format
    )


# =====================================================================
# 22. DASHBOARD HELPER TABLES
# =====================================================================
#
# Kept far to the right.
# Charts read these dynamic formulas.
# =====================================================================


# ---------------------------------------------------------------------
# A. Country count by Region   (cols X / Y)
# ---------------------------------------------------------------------

dash["X2"] = "Region"

dash["Y2"] = "Countries"


region_first_row = 3

region_last_row = (
    region_first_row
    + len(REGION_ORDER)
    - 1
)


for idx, region in enumerate(
    REGION_ORDER,
    start=region_first_row
):

    dash.cell(
        idx,
        24
    ).value = region


    dash.cell(
        idx,
        25
    ).value = (

        f'=IF('
        f'AND($C$4<>"All",$C$4<>X{idx}),0,'

        f'IF($C$7="All",'

        f'COUNTIF('
        f'GDP_DATA[UN Region],X{idx}),'

        f'COUNTIFS('
        f'GDP_DATA[UN Region],X{idx},'
        f'GDP_DATA[GDP Bucket],$C$7)))'
    )


# ---------------------------------------------------------------------
# B. Total GDP by Region       (cols AA / AB)
# ---------------------------------------------------------------------

dash["AA2"] = "Region"

dash["AB2"] = "Total GDP"


for idx, region in enumerate(
    REGION_ORDER,
    start=region_first_row
):

    dash.cell(
        idx,
        27
    ).value = region


    dash.cell(
        idx,
        28
    ).value = (

        f'=IF('
        f'AND($C$4<>"All",$C$4<>AA{idx}),0,'

        f'IF($C$7="All",'

        f'SUMIFS('
        f'GDP_DATA[IMF Estimate],'
        f'GDP_DATA[UN Region],AA{idx}),'

        f'SUMIFS('
        f'GDP_DATA[IMF Estimate],'
        f'GDP_DATA[UN Region],AA{idx},'
        f'GDP_DATA[GDP Bucket],$C$7)))'
    )


# ---------------------------------------------------------------------
# C. Average GDP by Region     (cols AD / AE)
# ---------------------------------------------------------------------

dash["AD2"] = "Region"

dash["AE2"] = "Average GDP"


for idx, region in enumerate(
    REGION_ORDER,
    start=region_first_row
):

    dash.cell(
        idx,
        30
    ).value = region


    dash.cell(
        idx,
        31
    ).value = (

        f'=IF('
        f'AND($C$4<>"All",$C$4<>AD{idx}),NA(),'

        f'IFERROR('

        f'IF($C$7="All",'

        f'AVERAGEIFS('
        f'GDP_DATA[IMF Estimate],'
        f'GDP_DATA[UN Region],AD{idx}),'

        f'AVERAGEIFS('
        f'GDP_DATA[IMF Estimate],'
        f'GDP_DATA[UN Region],AD{idx},'
        f'GDP_DATA[GDP Bucket],$C$7))'

        f',0))'
    )


# ---------------------------------------------------------------------
# D. Source comparison totals  (cols AG / AH)
# ---------------------------------------------------------------------

dash["AG2"] = "Source"

dash["AH2"] = "Total GDP"


source_first_row = 3

source_last_row = (
    source_first_row
    + len(SOURCE_NAMES)
    - 1
)


for idx, name in enumerate(
    SOURCE_NAMES,
    start=source_first_row
):

    dash.cell(
        idx,
        33
    ).value = name


    dash.cell(
        idx,
        34
    ).value = sum_formula(
        f"{name} Estimate"
    )


# ---------------------------------------------------------------------
# E. GDP Bucket distribution   (cols X / Y, lower block)
# ---------------------------------------------------------------------

dash["X15"] = "GDP Bucket"

dash["Y15"] = "Countries"


bucket_first_row = 16

bucket_last_row = (
    bucket_first_row
    + len(BUCKET_LABELS)
    - 1
)


for idx, label in enumerate(
    BUCKET_LABELS,
    start=bucket_first_row
):

    dash.cell(
        idx,
        24
    ).value = label


    dash.cell(
        idx,
        25
    ).value = (

        f'=IF('
        f'AND($C$7<>"All",$C$7<>X{idx}),0,'

        f'IF($C$4="All",'

        f'COUNTIF('
        f'GDP_DATA[GDP Bucket],X{idx}),'

        f'COUNTIFS('
        f'GDP_DATA[GDP Bucket],X{idx},'
        f'GDP_DATA[UN Region],$C$4)))'
    )


# ---------------------------------------------------------------------
# F. Top 10 economies          (cols AA / AB, lower block)
# ---------------------------------------------------------------------

dash["AA15"] = "Country"

dash["AB15"] = "IMF Estimate"


top_countries = (
    df
    .nlargest(
        10,
        "IMF Estimate"
    )
    ["Country"]
    .tolist()
)


top_first_row = 16

top_last_row = (
    top_first_row
    + len(top_countries)
    - 1
)


for idx, country in enumerate(
    top_countries,
    start=top_first_row
):

    dash.cell(
        idx,
        27
    ).value = country


    dash.cell(
        idx,
        28
    ).value = (

        f'=IFERROR('

        f'IF(AND($C$4="All",$C$7="All"),'

        f'SUMIFS('
        f'GDP_DATA[IMF Estimate],'
        f'GDP_DATA[Country],AA{idx}),'

        f'IF($C$4="All",'

        f'SUMIFS('
        f'GDP_DATA[IMF Estimate],'
        f'GDP_DATA[Country],AA{idx},'
        f'GDP_DATA[GDP Bucket],$C$7),'

        f'IF($C$7="All",'

        f'SUMIFS('
        f'GDP_DATA[IMF Estimate],'
        f'GDP_DATA[Country],AA{idx},'
        f'GDP_DATA[UN Region],$C$4),'

        f'SUMIFS('
        f'GDP_DATA[IMF Estimate],'
        f'GDP_DATA[Country],AA{idx},'
        f'GDP_DATA[UN Region],$C$4,'
        f'GDP_DATA[GDP Bucket],$C$7))))'

        f',0)'
    )


# =====================================================================
# 23. CREATE SIX DASHBOARD CHARTS
# =====================================================================

# ---------------------------------------------------------------------
# Chart 1 - Countries by Region
# ---------------------------------------------------------------------

chart1 = BarChart()

chart1.type = "bar"

chart1.style = 10

chart1.title = (
    "Countries by Region"
)

chart1.height = 7

chart1.width = 12


chart1.add_data(

    Reference(
        dash,
        min_col=25,
        min_row=2,
        max_row=region_last_row
    ),

    titles_from_data=True
)


chart1.set_categories(

    Reference(
        dash,
        min_col=24,
        min_row=region_first_row,
        max_row=region_last_row
    )
)


chart1.legend = None


dash.add_chart(
    chart1,
    "H2"
)


# ---------------------------------------------------------------------
# Chart 2 - Share of world GDP by Region
# ---------------------------------------------------------------------

chart2 = DoughnutChart()

chart2.title = (
    "Share of World GDP by Region"
)

chart2.height = 7

chart2.width = 10


chart2.add_data(

    Reference(
        dash,
        min_col=28,
        min_row=2,
        max_row=region_last_row
    ),

    titles_from_data=True
)


chart2.set_categories(

    Reference(
        dash,
        min_col=27,
        min_row=region_first_row,
        max_row=region_last_row
    )
)


dash.add_chart(
    chart2,
    "P2"
)


# ---------------------------------------------------------------------
# Chart 3 - Average GDP by Region
# ---------------------------------------------------------------------

chart3 = BarChart()

chart3.type = "col"

chart3.style = 11

chart3.title = (
    "Average GDP by Region (US$ mn)"
)

chart3.height = 7

chart3.width = 12


chart3.add_data(

    Reference(
        dash,
        min_col=31,
        min_row=2,
        max_row=region_last_row
    ),

    titles_from_data=True
)


chart3.set_categories(

    Reference(
        dash,
        min_col=30,
        min_row=region_first_row,
        max_row=region_last_row
    )
)


chart3.legend = None


dash.add_chart(
    chart3,
    "H17"
)


# ---------------------------------------------------------------------
# Chart 4 - IMF vs World Bank vs United Nations
# ---------------------------------------------------------------------

chart4 = BarChart()

chart4.type = "col"

chart4.style = 10

chart4.title = (
    "Total GDP by Source (US$ mn)"
)

chart4.height = 7

chart4.width = 10


chart4.add_data(

    Reference(
        dash,
        min_col=34,
        min_row=2,
        max_row=source_last_row
    ),

    titles_from_data=True
)


chart4.set_categories(

    Reference(
        dash,
        min_col=33,
        min_row=source_first_row,
        max_row=source_last_row
    )
)


chart4.legend = None


dash.add_chart(
    chart4,
    "P17"
)


# ---------------------------------------------------------------------
# Chart 5 - GDP size distribution
# ---------------------------------------------------------------------

chart5 = DoughnutChart()

chart5.title = (
    "Countries by GDP Size"
)

chart5.height = 7

chart5.width = 12


chart5.add_data(

    Reference(
        dash,
        min_col=25,
        min_row=15,
        max_row=bucket_last_row
    ),

    titles_from_data=True
)


chart5.set_categories(

    Reference(
        dash,
        min_col=24,
        min_row=bucket_first_row,
        max_row=bucket_last_row
    )
)


dash.add_chart(
    chart5,
    "H32"
)


# ---------------------------------------------------------------------
# Chart 6 - Top 10 economies
# ---------------------------------------------------------------------

chart6 = LineChart()

chart6.style = 13

chart6.title = (
    "Top 10 Economies (US$ mn)"
)

chart6.height = 7

chart6.width = 10


chart6.add_data(

    Reference(
        dash,
        min_col=28,
        min_row=15,
        max_row=top_last_row
    ),

    titles_from_data=True
)


chart6.set_categories(

    Reference(
        dash,
        min_col=27,
        min_row=top_first_row,
        max_row=top_last_row
    )
)


chart6.legend = None


dash.add_chart(
    chart6,
    "P32"
)


# =====================================================================
# 24. DASHBOARD COLUMN WIDTHS
# =====================================================================

dash.column_dimensions[
    "A"
].width = 22


dash.column_dimensions[
    "B"
].width = 22


dash.column_dimensions[
    "C"
].width = 25


dash.column_dimensions[
    "D"
].width = 22


dash.column_dimensions[
    "E"
].width = 20


dash.column_dimensions[
    "F"
].width = 18


for row in [
    4,
    7
]:

    dash.row_dimensions[
        row
    ].height = 28


# =====================================================================
# 25. VALIDATION SHEET FORMATTING
# =====================================================================

validation_ws.freeze_panes = (
    "A2"
)


for cell in validation_ws[1]:

    cell.font = Font(
        bold=True
    )


# =====================================================================
# 26. SAVE
# =====================================================================

wb.save(
    OUTPUT_FILE
)


# =====================================================================
# 27. VERIFY GDP_DATA EXCEL TABLE ROW COUNT
# =====================================================================

check_wb = load_workbook(
    OUTPUT_FILE
)


check_ws = check_wb[
    DATA_SHEET
]


if (
    EXCEL_TABLE_NAME
    not in check_ws.tables
):

    raise Exception(
        "Excel Table GDP_DATA was not created."
    )


check_table = check_ws.tables[
    EXCEL_TABLE_NAME
]


min_col, min_row, max_col, max_row = (
    range_boundaries(
        check_table.ref
    )
)


excel_rows = (
    max_row
    - min_row
)


check_wb.close()


# =====================================================================
# 28. FINAL COUNT VALIDATION
# =====================================================================

print(
    "\n"
    + "=" * 72
)

print(
    "FINAL WEBSITE / DATAFRAME / EXCEL VALIDATION"
)

print(
    "=" * 72
)


print(
    "Website total       :",
    website_total
)


print(
    "DataFrame rows      :",
    len(df)
)


print(
    "GDP_DATA Excel rows :",
    excel_rows
)


if not (

    website_total
    == len(df)
    == excel_rows

):

    raise Exception(

        "\n❌ FINAL COUNT MISMATCH\n\n"

        f"Website   = {website_total}\n"

        f"DataFrame = {len(df)}\n"

        f"Excel     = {excel_rows}\n\n"

        "Workbook will NOT be downloaded."
    )


print(
    "\n✅ WEBSITE TOTAL = DATAFRAME TOTAL = EXCEL TOTAL"
)


print(
    "\nWorkbook created with:"
)

print(
    "✓ GDP Data sheet"
)

print(
    "✓ Excel Table GDP_DATA"
)

print(
    "✓ Sorted by IMF estimate"
)

print(
    "✓ S.No. first column"
)

print(
    "✓ 6 Pivot Analysis tables"
)

print(
    "✓ Dashboard"
)

print(
    "✓ Region filter"
)

print(
    "✓ GDP size filter"
)

print(
    "✓ 6 dynamic dashboard charts"
)


print(
    "\nOutput:",
    OUTPUT_FILE
)

print(
    "CSV   :",
    CSV_FILE
)


# =====================================================================
# 29. DOWNLOAD
# =====================================================================

files.download(
    OUTPUT_FILE
)
